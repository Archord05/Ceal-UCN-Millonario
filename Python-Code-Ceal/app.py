"""
Quien Quiere Ser Millonario? - Aplicacion Web
Servidor Flask que lee preguntas de un archivo Excel y las sirve
a una interfaz web interactiva estilo el famoso programa de TV.

Formato Excel requerido:
  Pregunta | Opcion_A | Opcion_B | Opcion_C | Opcion_D | Respuesta_Correcta | Dificultad
  (Dificultad: 1=facil, 2=medio, 3=dificil)
"""

import json
import os
import random
from pathlib import Path

from dotenv import load_dotenv
from flask import Flask, jsonify, render_template, request
from openpyxl import Workbook, load_workbook

# Cargar variables de entorno desde .env
BASE_DIR = Path(__file__).parent
load_dotenv(BASE_DIR / ".env")

app = Flask(__name__)

# --- Configuracion desde .env ---
EXCEL_FILE = Path(os.getenv("EXCEL_FILE", "preguntas.xlsx"))
JSON_CACHE = Path(os.getenv("JSON_CACHE", "preguntas.json"))

# Si las rutas son relativas, resolverlas desde el directorio del proyecto
if not EXCEL_FILE.is_absolute():
    EXCEL_FILE = BASE_DIR / EXCEL_FILE
if not JSON_CACHE.is_absolute():
    JSON_CACHE = BASE_DIR / JSON_CACHE

PORT = int(os.getenv("PORT", "5000"))
HOST = os.getenv("HOST", "0.0.0.0")
DEBUG = os.getenv("DEBUG", "True").lower() in ("true", "1", "yes")

# Premios: Puedes editar esta lista manualmente para cambiar los premios de cada pregunta.
# Asegurate de que haya exactamente 15 premios en la lista.
PRIZE_LEVELS = [
    "150 pts",
    "300 pts",
    "450 pts",
    "600 pts",
    "750 pts",
    "900 pts",
    "1050 pts",
    "1200 pts",
    "1350 pts",
    "1500 pts",
    "1650 pts",
    "1800 pts",
    "1950 pts",
    "2100 pts",
    "Premio Sorpresa Grande"
]

# Distribucion de dificultad por pregunta (5 preguntas por nivel)
# Preguntas 1-5: dificultad 1 (facil), 6-10: dificultad 2 (medio), 11-15: dificultad 3 (dificil)
DIFFICULTY_DISTRIBUTION = [1, 1, 1, 1, 1, 2, 2, 2, 2, 2, 3, 3, 3, 3, 3]
QUESTIONS_PER_GAME = 15


def create_sample_excel():
    """Crea un archivo Excel de ejemplo con preguntas y dificultad si no existe."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Preguntas"

    # Encabezados (7 columnas)
    headers = [
        "Pregunta", "Opcion_A", "Opcion_B", "Opcion_C",
        "Opcion_D", "Respuesta_Correcta", "Dificultad"
    ]
    ws.append(headers)

    # Preguntas de ejemplo organizadas por dificultad (1=facil, 2=medio, 3=dificil)
    sample_questions = [
        # --- DIFICULTAD 1 (Facil) ---
        ["Cuantos dias tiene una semana?", "5", "6", "7", "8", "C", 1],
        ["De que color es el cielo en un dia despejado?", "Verde", "Rojo", "Azul", "Amarillo", "C", 1],
        ["Cuantas patas tiene un perro?", "2", "4", "6", "8", "B", 1],
        ["Que animal hace 'muu'?", "Gato", "Perro", "Vaca", "Gallina", "C", 1],
        ["Cual es el resultado de 2 + 2?", "3", "4", "5", "6", "B", 1],
        ["En que estacion del anio hace mas calor?", "Invierno", "Otonio", "Primavera", "Verano", "D", 1],
        ["Cual es el oceano mas grande del mundo?", "Pacifico", "Atlantico", "Indico", "Artico", "A", 1],
        ["Cuantos meses tiene un anio?", "10", "11", "12", "13", "C", 1],
        ["Que animal es conocido como el rey de la selva?", "Tigre", "Elefante", "Leon", "Oso", "C", 1],
        ["En que continente esta Brasil?", "Europa", "Asia", "Africa", "America del Sur", "D", 1],

        # --- DIFICULTAD 2 (Medio) ---
        ["Cual es el planeta mas cercano al Sol?", "Venus", "Mercurio", "Marte", "Tierra", "B", 2],
        ["Quien pinto la Mona Lisa?", "Picasso", "Da Vinci", "Miguel Angel", "Rafael", "B", 2],
        ["Cual es el elemento quimico con simbolo 'O'?", "Oro", "Osmio", "Oxigeno", "Oganesson", "C", 2],
        ["En que anio llego el hombre a la Luna?", "1965", "1970", "1969", "1972", "C", 2],
        ["Cual es la capital de Australia?", "Sidney", "Melbourne", "Canberra", "Perth", "C", 2],
        ["Quien escribio 'Cien Anios de Soledad'?", "Borges", "Neruda", "Vargas Llosa", "Garcia Marquez", "D", 2],
        ["Cual es el rio mas largo del mundo?", "Nilo", "Amazonas", "Misisipi", "Yangtsee", "B", 2],
        ["Quien descubrio la penicilina?", "Pasteur", "Fleming", "Koch", "Jenner", "B", 2],
        ["Cual es la moneda oficial de Japon?", "Yuan", "Won", "Yen", "Ringgit", "C", 2],
        ["En que anio se fundo la ONU?", "1940", "1945", "1950", "1948", "B", 2],

        # --- DIFICULTAD 3 (Dificil) ---
        ["Cual es la velocidad de la luz en km/s (aprox)?", "200.000", "350.000", "300.000", "150.000", "C", 3],
        ["Cuantos huesos tiene el cuerpo humano adulto?", "186", "206", "226", "196", "B", 3],
        ["Que elemento tiene el numero atomico 79?", "Plata", "Platino", "Oro", "Cobre", "C", 3],
        ["En que anio se publico 'El Quijote'?", "1505", "1605", "1705", "1580", "B", 3],
        ["Cual es el planeta con mas lunas en el sistema solar?", "Jupiter", "Saturno", "Urano", "Neptuno", "B", 3],
        ["Que teorema establece que a^2 + b^2 = c^2?", "Fermat", "Euler", "Pitagoras", "Tales", "C", 3],
        ["Cual es el hueso mas pequenio del cuerpo humano?", "Estribo", "Martillo", "Yunque", "Femur", "A", 3],
        ["Que gas es mas abundante en la atmosfera?", "Oxigeno", "Hidrogeno", "Nitrogeno", "CO2", "C", 3],
        ["En que anio cayo el Muro de Berlin?", "1987", "1989", "1991", "1990", "B", 3],
        ["Cual es el metal mas abundante en la corteza terrestre?", "Hierro", "Aluminio", "Cobre", "Zinc", "B", 3],
    ]

    for q in sample_questions:
        ws.append(q)

    # Ajustar anchos de columna
    ws.column_dimensions["A"].width = 55
    for col in ["B", "C", "D", "E"]:
        ws.column_dimensions[col].width = 20
    ws.column_dimensions["F"].width = 22
    ws.column_dimensions["G"].width = 12

    wb.save(EXCEL_FILE)
    print(f"[OK] Archivo de preguntas creado: {EXCEL_FILE}")


def load_questions_from_excel():
    """
    Lee las preguntas del archivo Excel y las guarda en un JSON de cache.
    Retorna la lista de preguntas como diccionarios, agrupadas por dificultad.
    """
    if not EXCEL_FILE.exists():
        create_sample_excel()

    wb = load_workbook(EXCEL_FILE, read_only=True)
    ws = wb.active

    questions = []
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=1):
        # Saltar filas vacias
        if not row or not row[0]:
            continue

        # Leer 7 columnas (con dificultad)
        values = list(row[:7])
        # Rellenar columnas faltantes
        while len(values) < 7:
            values.append(None)

        pregunta, op_a, op_b, op_c, op_d, correcta, dificultad = values

        # Mapear la letra de la respuesta correcta al indice
        answer_map = {"A": 0, "B": 1, "C": 2, "D": 3}
        correct_index = answer_map.get(str(correcta).strip().upper(), 0)

        # Parsear dificultad (default 2 si no se especifica)
        try:
            diff = int(dificultad) if dificultad is not None else 2
            diff = max(1, min(3, diff))  # Clamp entre 1 y 3
        except (ValueError, TypeError):
            diff = 2

        questions.append({
            "id": row_idx,
            "question": str(pregunta).strip(),
            "options": [
                str(op_a).strip() if op_a else "",
                str(op_b).strip() if op_b else "",
                str(op_c).strip() if op_c else "",
                str(op_d).strip() if op_d else "",
            ],
            "correct": correct_index,
            "difficulty": diff,
        })

    wb.close()

    # Guardar en cache JSON
    with open(JSON_CACHE, "w", encoding="utf-8") as f:
        json.dump(questions, f, ensure_ascii=False, indent=2)

    # Resumen por dificultad
    diff_counts = {}
    for q in questions:
        d = q["difficulty"]
        diff_counts[d] = diff_counts.get(d, 0) + 1

    print(f"[INFO] {len(questions)} preguntas cargadas desde Excel -> {JSON_CACHE}")
    for d in sorted(diff_counts):
        print(f"       Dificultad {d}: {diff_counts[d]} preguntas")

    return questions


def get_questions():
    """Obtiene las preguntas, primero del cache JSON, luego del Excel."""
    if JSON_CACHE.exists():
        # Verificar si el Excel es mas reciente que el cache
        if EXCEL_FILE.exists():
            excel_mtime = os.path.getmtime(EXCEL_FILE)
            json_mtime = os.path.getmtime(JSON_CACHE)
            if excel_mtime > json_mtime:
                return load_questions_from_excel()

        with open(JSON_CACHE, "r", encoding="utf-8") as f:
            return json.load(f)

    return load_questions_from_excel()


def select_game_questions(all_questions):
    """
    Selecciona 15 preguntas al azar respetando la dificultad creciente.
    Filtra las ya usadas en la sesion.
    """
    # Filtrar las no usadas
    available_questions = [q for q in all_questions if not q.get("used", False)]

    # Si no nos alcanzan para un juego (15), reseteamos el historial
    if len(available_questions) < QUESTIONS_PER_GAME:
        for q in all_questions:
            q["used"] = False
        available_questions = all_questions
        print("[INFO] No quedan suficientes preguntas nuevas. Reseteando historial de uso.")
        # Guardar el reseteo en el JSON
        try:
            with open(JSON_CACHE, "w", encoding="utf-8") as f:
                json.dump(all_questions, f, ensure_ascii=False, indent=2)
        except Exception as e:
            print(f"[ERROR] No se pudo actualizar el historial en JSON: {e}")

    # Agrupar por dificultad
    by_difficulty = {1: [], 2: [], 3: []}
    for q in available_questions:
        d = q.get("difficulty", 2)
        if d in by_difficulty:
            by_difficulty[d].append(q)

    selected = []
    for difficulty_level in DIFFICULTY_DISTRIBUTION:
        pool = by_difficulty.get(difficulty_level, [])
        available_in_pool = [q for q in pool if q not in selected]

        if available_in_pool:
            chosen = random.choice(available_in_pool)
            selected.append(chosen)
        else:
            # Buscar en niveles adyacentes si no hay suficientes
            for offset in [1, -1, 2, -2]:
                alt_level = difficulty_level + offset
                alt_pool = by_difficulty.get(alt_level, [])
                alt_available = [q for q in alt_pool if q not in selected]
                if alt_available:
                    chosen = random.choice(alt_available)
                    selected.append(chosen)
                    break

    return selected


# --- Rutas ---

@app.route("/")
def index():
    """Pagina principal del juego."""
    import time
    return render_template("index.html", t=int(time.time()))


@app.route("/api/game")
def api_game():
    """
    Devuelve un set de 15 preguntas randomizadas con dificultad creciente
    junto con los premios. Cada llamada genera una partida diferente.
    """
    all_questions = get_questions()
    game_questions = select_game_questions(all_questions)

    # Preparamos las preguntas para el frontend
    clean_questions = []
    for q in game_questions:
        clean_questions.append({
            "id": q["id"],
            "question": q["question"],
            "options": q["options"],
            "correct": q["correct"],
            "difficulty": q["difficulty"]
        })

    return jsonify({
        "questions": clean_questions,
        "prizes": PRIZE_LEVELS,
        "total_pool": len(all_questions),
    })


@app.route("/api/reload", methods=["POST"])
def api_reload():
    """Fuerza la recarga de preguntas desde el archivo Excel."""
    questions = load_questions_from_excel()

    diff_counts = {}
    for q in questions:
        d = q["difficulty"]
        diff_counts[d] = diff_counts.get(d, 0) + 1

    return jsonify({
        "status": "ok",
        "count": len(questions),
        "by_difficulty": diff_counts,
    })


@app.route("/api/mark_used", methods=["POST"])
def api_mark_used():
    """Marca una pregunta individual como usada en el JSON."""
    data = request.json
    question_id = data.get("id")
    
    if not question_id:
        return jsonify({"status": "error"}), 400

    if JSON_CACHE.exists():
        try:
            with open(JSON_CACHE, "r", encoding="utf-8") as f:
                all_questions = json.load(f)
                
            for q in all_questions:
                if q["id"] == question_id:
                    q["used"] = True
                    break
                    
            with open(JSON_CACHE, "w", encoding="utf-8") as f:
                json.dump(all_questions, f, ensure_ascii=False, indent=2)
        except Exception as e:
            print(f"[ERROR] Error al marcar pregunta usada: {e}")
            
    return jsonify({"status": "ok"})


@app.route("/api/debug", methods=["POST"])
def api_debug():
    """Recibe la informacion de la pregunta actual del frontend y la imprime en la consola de Python."""
    data = request.json
    print(f"\n" + "="*50)
    print(f"--- DEBUG: PREGUNTA {data.get('current')} DE {data.get('total')} ---")
    print(f"Fila Excel: {data.get('id')} | Dificultad: Nivel {data.get('difficulty')}")
    print(f"Pregunta:   {data.get('question')}")
    opts = data.get('options', [])
    print(f"Opciones:   A) {opts[0]} | B) {opts[1]} | C) {opts[2]} | D) {opts[3]}")
    print(f"Respuesta:  {data.get('correct_letter')} (Indice {data.get('correct_index')})")
    print("="*50 + "\n")
    return jsonify({"status": "ok"})


if __name__ == "__main__":
    # Cargar preguntas al inicio
    load_questions_from_excel()
    print(f"\n*** Quien Quiere Ser Millonario! ***")
    print(f">>> Abre tu navegador en: http://localhost:{PORT}\n")
    print(f"    Excel: {EXCEL_FILE}")
    print(f"    Cache: {JSON_CACHE}\n")
    app.run(debug=DEBUG, host=HOST, port=PORT)
